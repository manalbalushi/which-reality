import io

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook

from ..db import get_db
from .. import utils

bp = Blueprint("register", __name__)


def _build_query(args):
    clauses = []
    params = []

    def like(field, value):
        clauses.append(f"{field} LIKE ?")
        params.append(f"%{value}%")

    if args.get("q"):
        clauses.append("(action_id LIKE ? OR title LIKE ? OR description LIKE ? OR latest_comment LIKE ?)")
        q = f"%{args['q']}%"
        params.extend([q, q, q, q])
    if args.get("owner"):
        like("owner", args["owner"])
    if args.get("department"):
        like("department", args["department"])
    if args.get("status"):
        clauses.append("status = ?")
        params.append(args["status"])
    if args.get("priority"):
        clauses.append("priority = ?")
        params.append(args["priority"])
    if args.get("source"):
        like("source", args["source"])
    if args.get("evidence_status"):
        clauses.append("evidence_status = ?")
        params.append(args["evidence_status"])
    if args.get("target_date_from"):
        clauses.append("current_target_date >= ?")
        params.append(args["target_date_from"])
    if args.get("target_date_to"):
        clauses.append("current_target_date <= ?")
        params.append(args["target_date_to"])

    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    return where, params


@bp.route("/register")
@utils.login_required
def register():
    db = get_db()
    where, params = _build_query(request.args)
    rows = db.execute(f"SELECT * FROM actions {where} ORDER BY action_id", params).fetchall()
    actions = []
    for r in rows:
        d = dict(r)
        d["days_open"] = utils.days_open(r)
        d["days_overdue"] = utils.days_overdue(r)
        actions.append(d)

    owners = [r["name"] for r in db.execute("SELECT DISTINCT name FROM users ORDER BY name")]
    departments = [r["department"] for r in db.execute(
        "SELECT DISTINCT department FROM actions WHERE department IS NOT NULL ORDER BY department"
    )]

    return render_template(
        "register.html", actions=actions, statuses=utils.STATUS_OPTIONS,
        priorities=utils.PRIORITY_OPTIONS, evidence_statuses=utils.EVIDENCE_STATUS_OPTIONS,
        owners=owners, departments=departments, args=request.args,
    )


@bp.route("/register/export")
@utils.login_required
def export():
    db = get_db()
    where, params = _build_query(request.args)
    rows = db.execute(f"SELECT * FROM actions {where} ORDER BY action_id", params).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Action Register"
    headers = [
        "Action ID", "Title", "Description", "Owner", "Department", "Priority", "Source",
        "Date Raised", "Original Target Date", "Current Target Date", "Status", "Completion %",
        "Latest Comment", "Last Updated", "Days Open", "Days Overdue", "Evidence Required",
        "Evidence Status", "Closure Date", "Legacy ID",
    ]
    ws.append(headers)
    for r in rows:
        ws.append([
            r["action_id"], r["title"], r["description"], r["owner"], r["department"],
            r["priority"], r["source"], utils.dateonly(r["date_raised"]),
            utils.dateonly(r["original_target_date"]), utils.dateonly(r["current_target_date"]),
            r["status"], r["completion_pct"], r["latest_comment"], utils.dateonly(r["last_updated"]),
            utils.days_open(r), utils.days_overdue(r), r["evidence_required"], r["evidence_status"],
            utils.dateonly(r["closure_date"]), r["legacy_id"],
        ])
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 50)

    history_ws = wb.create_sheet("Update History")
    history_ws.append([
        "Action ID", "Update Date", "Updated By", "Previous Status", "New Status",
        "Previous Target Date", "New Target Date", "Comment", "Evidence Submitted", "Approval Status",
    ])
    for r in db.execute("SELECT * FROM update_history ORDER BY action_id, update_date"):
        history_ws.append([
            r["action_id"], utils.dateonly(r["update_date"]), r["updated_by"], r["previous_status"],
            r["new_status"], utils.dateonly(r["previous_target_date"]), utils.dateonly(r["new_target_date"]),
            r["comment"], r["evidence_submitted"], r["approval_status"],
        ])

    evidence_ws = wb.create_sheet("Evidence")
    evidence_ws.append([
        "Evidence ID", "Action ID", "Evidence Name", "Type", "Description", "Uploaded By",
        "Upload Date", "Version", "Evidence Status", "Reviewer", "Review Date", "Rejection Reason",
    ])
    for r in db.execute("SELECT * FROM evidence ORDER BY action_id, evidence_id"):
        evidence_ws.append([
            r["evidence_id"], r["action_id"], r["evidence_name"], r["evidence_type"], r["description"],
            r["uploaded_by"], utils.dateonly(r["upload_date"]), r["version"], r["evidence_status"],
            r["reviewer"], utils.dateonly(r["review_date"]), r["rejection_reason"],
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"action_register_export_{utils.today_str()}.xlsx"
    return send_file(
        buf, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
