import os
import re
from datetime import datetime, date

from flask import (
    Blueprint, render_template, request, redirect, url_for, flash, current_app,
)
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from ..db import get_db
from .. import utils

bp = Blueprint("importer", __name__)

TARGET_FIELDS = [
    ("title", "Action Title", True),
    ("description", "Action Description", False),
    ("owner", "Action Owner", True),
    ("department", "Department / Function", False),
    ("priority", "Priority", False),
    ("source", "Source / Assessment", False),
    ("date_raised", "Date Raised", False),
    ("original_target_date", "Original Target Date", False),
    ("current_target_date", "Current Target Date", False),
    ("status", "Status", False),
    ("completion_pct", "Completion %", False),
    ("latest_comment", "Latest Comment / Progress", False),
    ("evidence_required", "Evidence Required (Yes/No)", False),
    ("legacy_id", "Original Action Reference / ID (kept, not overwritten)", False),
]

_ALIASES = {
    "title": ["action", "action title", "title", "action item", "finding"],
    "description": ["description", "detail", "details", "action description"],
    "owner": ["owner", "action owner", "assigned to", "responsible"],
    "department": ["department", "function", "team", "business unit", "dept"],
    "priority": ["priority", "risk rating", "severity"],
    "source": ["source", "assessment", "audit", "origin"],
    "date_raised": ["date raised", "raised date", "identified date", "opened date"],
    "original_target_date": ["original target date", "original due date"],
    "current_target_date": ["target date", "due date", "current target date"],
    "status": ["status"],
    "completion_pct": ["completion", "% complete", "progress", "completion %"],
    "latest_comment": ["comment", "comments", "notes", "progress update", "remarks"],
    "evidence_required": ["evidence required", "evidence needed"],
    "legacy_id": ["ref", "ref no", "id", "action id", "reference", "action ref"],
}


def _normalize(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def _auto_map(headers):
    mapping = {}
    used = set()
    normalized_headers = [_normalize(h) for h in headers]
    for field, _, _ in TARGET_FIELDS:
        best_idx = None
        for alias in _ALIASES.get(field, []) + [field]:
            alias_n = _normalize(alias)
            for i, h in enumerate(normalized_headers):
                if i in used:
                    continue
                if h == alias_n or (alias_n and alias_n in h):
                    best_idx = i
                    break
            if best_idx is not None:
                break
        if best_idx is not None:
            mapping[field] = best_idx
            used.add(best_idx)
    return mapping


def _cell_str(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


@bp.route("/import", methods=["GET"])
@utils.role_required("Admin")
def import_home():
    return render_template("import_upload.html")


@bp.route("/import/upload", methods=["POST"])
@utils.role_required("Admin")
def import_upload():
    file = request.files.get("excel_file")
    if not file or not file.filename:
        flash("Please choose an Excel (.xlsx) file to upload.", "danger")
        return redirect(url_for("importer.import_home"))
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        flash("Only .xlsx/.xlsm files are supported.", "danger")
        return redirect(url_for("importer.import_home"))

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    stored_name = f"{timestamp}_{secure_filename(file.filename)}"
    folder = os.path.join(current_app.config["UPLOAD_FOLDER"], "imports")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, stored_name)
    file.save(path)  # original file preserved untouched, never edited or deleted

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h) if h is not None else f"Column {i+1}" for i, h in enumerate(next(rows_iter))]
    except StopIteration:
        flash("That file appears to be empty.", "danger")
        return redirect(url_for("importer.import_home"))

    preview_rows = []
    for i, row in enumerate(rows_iter):
        if i >= 5:
            break
        preview_rows.append([_cell_str(v) for v in row])

    mapping = _auto_map(headers)

    return render_template(
        "import_mapping.html", stored_name=stored_name, headers=headers,
        target_fields=TARGET_FIELDS, mapping=mapping, preview_rows=preview_rows,
    )


@bp.route("/import/confirm", methods=["POST"])
@utils.role_required("Admin")
def import_confirm():
    stored_name = request.form.get("stored_name")
    if not stored_name:
        flash("Import session expired, please upload the file again.", "danger")
        return redirect(url_for("importer.import_home"))

    path = os.path.join(current_app.config["UPLOAD_FOLDER"], "imports", stored_name)
    if not os.path.exists(path):
        flash("Uploaded file could not be found, please upload again.", "danger")
        return redirect(url_for("importer.import_home"))

    mapping = {}
    for field, _, _ in TARGET_FIELDS:
        idx = request.form.get(f"map_{field}", "")
        if idx != "":
            mapping[field] = int(idx)

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter)  # skip header

    db = get_db()
    imported = 0
    defaulted_rows = []
    now = utils.now_iso()
    today_iso = utils.today_str()

    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue

        def get(field, default=""):
            idx = mapping.get(field)
            if idx is None or idx >= len(row):
                return default
            return _cell_str(row[idx]) or default

        title = get("title")
        if not title:
            continue  # no title, no reliable action to import - skipped, original file untouched

        owner = get("owner", "Unassigned")
        defaults_used = []
        priority = get("priority", "Medium")
        if priority not in utils.PRIORITY_OPTIONS:
            priority = "Medium"
            defaults_used.append("priority")
        status = get("status", "Not Started")
        if status not in utils.STATUS_OPTIONS:
            status = "Not Started"
            defaults_used.append("status")
        date_raised = get("date_raised") or today_iso
        current_target_date = get("current_target_date") or get("original_target_date") or ""
        original_target_date = get("original_target_date") or current_target_date
        evidence_required = get("evidence_required", "Yes")
        if evidence_required not in utils.YES_NO:
            evidence_required = "Yes"
        completion_raw = get("completion_pct", "0")
        try:
            completion_pct = int(float(str(completion_raw).replace("%", "").strip() or 0))
            completion_pct = max(0, min(100, completion_pct))
        except ValueError:
            completion_pct = 0
            defaults_used.append("completion_pct")

        action_id = utils.next_action_id(db)
        db.execute(
            """INSERT INTO actions (
                action_id, title, description, owner, department, priority, source,
                date_raised, original_target_date, current_target_date, status,
                completion_pct, latest_comment, last_updated, evidence_required,
                evidence_status, closure_date, legacy_id, created_at, created_by
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                action_id, title, get("description"), owner, get("department"), priority,
                get("source"), date_raised, original_target_date or None, current_target_date or None,
                status, completion_pct, get("latest_comment"), now, evidence_required,
                "Not Submitted", None, get("legacy_id") or None, now, f"Excel import ({utils.current_user()})",
            ),
        )
        db.execute(
            """INSERT INTO update_history (
                action_id, update_date, updated_by, previous_status, new_status,
                previous_target_date, new_target_date, comment, evidence_submitted, approval_status
            ) VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (
                action_id, now, utils.current_user(), None, status, None, current_target_date or None,
                f"Imported from Excel tracker ({stored_name}), source row {row_num}.", "No", None,
            ),
        )
        utils.notify(db, owner, "Action assigned", action_id,
                     f"{action_id} \"{title}\" was assigned to you (migrated from Excel tracker).")
        imported += 1
        if defaults_used:
            defaulted_rows.append((row_num, action_id, defaults_used))

    utils.log_audit(db, utils.current_user(), "excel_import", None,
                     f"Imported {imported} action(s) from {stored_name}.")
    db.commit()

    flash(f"Import complete: {imported} action(s) created. Original file preserved for reference.", "success")
    return render_template("import_result.html", imported=imported, defaulted_rows=defaulted_rows, stored_name=stored_name)
